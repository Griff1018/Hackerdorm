"""Export the database to an .xlsx workbook (backup / spreadsheet view for the agent)."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from . import db
from .services.common import fmt_local

_TIME_COLS = {"created_at", "updated_at", "qualified_at", "followup_anchor_at", "next_followup_at",
              "last_contacted_at", "received_at", "processed_at"}


def _add_sheet(wb: Workbook, title: str, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(title)
    if not rows:
        return
    cols = list(rows[0].keys())
    ws.append(cols)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([fmt_local(v) if c in _TIME_COLS and v else v
                   for c, v in ((col, row.get(col)) for col in cols)])
    for i, col in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(len(col) + 2, 12), 40)
    ws.freeze_panes = "A2"


def build_workbook() -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    leads = [dict(r) for r in db.table_rows("leads")]
    lead_names = {r["id"]: r.get("name") for r in leads}
    messages = [dict(r) for r in db.table_rows("messages", "id, lead_id, role, channel, kind, content, created_at")]
    msg_rows = [{"id": m["id"], "lead_id": m["lead_id"], "lead_name": lead_names.get(m["lead_id"], ""),
                 "role": m["role"], "channel": m["channel"], "kind": m["kind"],
                 "content": m["content"], "created_at": m["created_at"]} for m in messages]
    _add_sheet(wb, "Leads", leads)
    _add_sheet(wb, "Messages", msg_rows)
    _add_sheet(wb, "Enquiries", [dict(r) for r in db.table_rows("enquiries")])
    _add_sheet(wb, "Listings", [dict(r) for r in db.table_rows("listings")])
    _add_sheet(wb, "Agent Profile", [dict(r) for r in db.table_rows("agent_profile")])
    _add_sheet(wb, "Settings", [dict(r) for r in db.table_rows("settings")])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

