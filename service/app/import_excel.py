"""One-off import of the legacy Google Sheets/Excel database into SQLite, with cleaning.

    python -m app.import_excel "../(FINAL) ULTIMATE Template - Ultimate Lead Database.xlsx"
    python -m app.import_excel <file.xlsx> --with-leads      # also import (deduplicated) leads

Legacy sheet layout (by column position, headers are unreliable):
  Sheet1 = leads, Sheet2 = listings, Sheet3 = agent profile, "Config Sheet" A2 = re-engagement ON/OFF.
Known corruption handled here: mixed date formats, numbers stored as floats, 'PENDING'
placeholders, status text in date columns, one lead duplicated across many rows,
inconsistent casing ('BUYER' vs 'Buyer'), "YES '+61..." phone cells.
"""
from __future__ import annotations

import argparse
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from . import db
from .config import settings
from .services import common

PLACEHOLDERS = {"", "pending", "none", "null", "n/a"}
STATUS_MAP = {"initiated": "INITIATED", "in progress": "IN_PROGRESS", "qualified": "QUALIFIED",
              "needs agent": "NEEDS_AGENT", "banned": "BANNED"}
DATE_FORMATS = ["%d %m %Y %H:%M", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y", "%d.%m.%Y %H:%M",
                "%d.%m.%Y", "%d %m %Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"]


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def clean(value: Any) -> str | None:
    t = text(value)
    return None if t.lower() in PLACEHOLDERS else t


def parse_date(value: Any) -> str | None:
    """Parse a legacy date cell as local time; reject unparseable or future values."""
    if value is None or value == "":
        return None
    dt: datetime | None = None
    if isinstance(value, datetime):
        dt = value
    else:
        raw = text(value)
        for fmt in DATE_FORMATS:
            try:
                dt = datetime.strptime(raw, fmt)
                break
            except ValueError:
                continue
    if dt is None:
        return None
    local = dt.replace(tzinfo=settings.tz) if dt.tzinfo is None else dt
    if local > common.local_now():
        return None  # day/month-swapped or otherwise bogus
    return db.to_iso(local)


def money(value: Any) -> str:
    if isinstance(value, (int, float)):
        return f"${value:,.0f}"
    return text(value)


def rows(ws: Any) -> list[tuple[Any, ...]]:
    return [r for r in ws.iter_rows(min_row=2, values_only=True) if any(v not in (None, "") for v in r)]


def cell(row: tuple[Any, ...], i: int) -> Any:
    return row[i] if i < len(row) else None


# --- importers -----------------------------------------------------------------------

def import_agent(wb: Any) -> str:
    if "Sheet3" not in wb.sheetnames or not rows(wb["Sheet3"]):
        return "agent profile: none found"
    r = rows(wb["Sheet3"])[0]
    phone_cell = text(cell(r, 4))
    phone_match = re.search(r"\+?\d[\d\s-]{5,}\d", phone_cell)
    db.save_agent_profile({
        "agent_name": text(cell(r, 0)), "agency": text(cell(r, 1)), "service_areas": text(cell(r, 2)),
        "working_hours": text(cell(r, 3)),
        "phone_number": phone_match.group(0).strip() if phone_match else "",
        "include_phone": 1 if phone_cell.upper().startswith("YES") else 0,
        "commission_policy": text(cell(r, 5)), "marketing_services": text(cell(r, 6)),
        "mortgage_partner": text(cell(r, 7)), "conveyancer_partner": text(cell(r, 8)),
        "showing_style": text(cell(r, 9)),
    })
    return f"agent profile: {text(cell(r, 0))}"


def import_settings(wb: Any) -> str:
    if "Config Sheet" not in wb.sheetnames:
        return "settings: none"
    value = text(wb["Config Sheet"]["A2"].value).upper()
    db.set_setting("reengagement_enabled", "1" if value == "ON" else "0")
    return f"re-engagement: {'ON' if value == 'ON' else 'OFF'}"


def import_listings(wb: Any) -> str:
    if "Sheet2" not in wb.sheetnames:
        return "listings: none"
    added = skipped = 0
    existing = db.listing_addresses()
    for r in rows(wb["Sheet2"]):
        address = text(cell(r, 0))
        if not address or address.lower() in existing:
            skipped += 1
            continue
        values = {f: text(cell(r, i)) for i, f in enumerate(db.LISTING_FIELDS)}
        values["asking_price"] = money(cell(r, 2))
        sold = text(cell(r, 25)).upper()
        now = db.now_iso()
        db.insert("listings", {
            **values,
            "contract_url": clean(cell(r, 23)),
            "status": "SOLD" if sold == "SOLD" else "SELLING",
            "review_token": db.new_review_token(),
            "created_at": now, "updated_at": now,
        })
        existing.add(address.lower())
        added += 1
    return f"listings: {added} added, {skipped} skipped"


_LINE_RE = re.compile(r"^(ORI Message|User|AI|Agent)\s*:\s*", re.IGNORECASE)


def split_history(history: str) -> list[tuple[str, str, str]]:
    """Legacy Chat_History blob -> [(role, kind, content)]."""
    out: list[tuple[str, str, str]] = []
    for line in (history or "").replace("\r\n", "\n").split("\n"):
        m = _LINE_RE.match(line.strip())
        if m:
            label = m.group(1).lower()
            role = {"user": "lead", "ori message": "lead", "ai": "ai", "agent": "agent"}[label]
            out.append((role, "enquiry" if label == "ori message" else "chat", line.strip()[m.end():].strip()))
        elif out and line.strip():
            role, kind, content = out[-1]
            out[-1] = (role, kind, f"{content}\n{line.strip()}")
    return out


def import_leads(wb: Any) -> str:
    if "Sheet1" not in wb.sheetnames:
        return "leads: none"
    groups: dict[str, list[tuple[Any, ...]]] = {}
    for r in rows(wb["Sheet1"]):
        key = (clean(cell(r, 1)) or "").lower() or common.normalize_phone(text(cell(r, 2))) or f"row-{id(r)}"
        groups.setdefault(key, []).append(r)

    added = skipped = 0
    for group in groups.values():
        latest = max(group, key=lambda r: parse_date(cell(r, 17)) or "")
        email = clean(cell(latest, 1))
        contact = text(cell(latest, 2)).lstrip("'")
        phone = common.normalize_phone(contact) if contact.startswith("+") else None
        chat_id = contact if contact.isdigit() and not phone else None
        if db.find_lead_by_contact(email, phone):
            skipped += 1
            continue
        status = STATUS_MAP.get(text(cell(latest, 5)).lower(), "INITIATED")
        score_raw = clean(cell(latest, 10))
        try:
            score = int(float(score_raw)) if score_raw else None
        except ValueError:
            score = None
        qualified = clean(cell(latest, 12))
        last_contacted = parse_date(cell(latest, 17))
        step = int(float(text(cell(latest, 16)) or 0)) if text(cell(latest, 16)).replace(".", "").isdigit() else 0
        active = status in ("INITIATED", "IN_PROGRESS")
        lead = db.create_lead({
            "name": text(cell(latest, 0)), "email": email, "phone": phone, "telegram_chat_id": chat_id,
            "lead_type": common.normalize_lead_type(text(cell(latest, 3))),
            "property_address": text(cell(latest, 4)),
            "listing_id": (db.match_listing(text(cell(latest, 4))) or {"id": None})["id"],
            "source": "IMPORT", "status": status, "is_returning": 1 if len(group) > 1 else 0,
            "timeline": clean(cell(latest, 7)), "budget": clean(cell(latest, 8)),
            "pre_approval": clean(cell(latest, 9)), "score": score, "score_reason": clean(cell(latest, 11)),
            "is_qualified": None if qualified is None else (1 if qualified.upper() == "YES" else 0),
            "notes": clean(cell(latest, 13)),
            "qualified_at": last_contacted if status == "QUALIFIED" else None,
            "followup_step": step if active else 0,
            "followup_anchor_at": last_contacted if active else None,
            "next_followup_at": common.next_followup_at(last_contacted, step) if active else None,
            "last_contacted_at": last_contacted,
        })
        for role, kind, content in split_history(text(cell(latest, 6))):
            db.insert("messages", {"lead_id": lead["id"], "role": role, "channel": "telegram", "kind": kind,
                                   "content": content, "created_at": last_contacted or db.now_iso()})
        added += 1
    return f"leads: {added} added from {sum(len(g) for g in groups.values())} rows, {skipped} already present"


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("--with-leads", action="store_true", help="also import leads (they are test data in the template)")
    args = parser.parse_args(argv)
    if not args.xlsx.exists():
        print(f"File not found: {args.xlsx}", file=sys.stderr)
        return 1
    db.connect()
    wb = load_workbook(args.xlsx, data_only=True)
    results = [import_agent(wb), import_settings(wb), import_listings(wb)]
    if args.with_leads:
        results.append(import_leads(wb))
    print(f"Imported into {settings.db_path}:")
    for line in results:
        print(" -", line)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
